# Gera base FICTÍCIA de rotas de entrega (last mile) no modelo Mercado Livre, para painel de demonstração.
# Indicadores genéricos (no prazo, insucesso, atraso de saída/retorno): confirmar os oficiais com o cliente.
# Mesma história das outras bases: jul-ago antes do serviço, set implantação, out com controle.
# Uso: python3 gerar_base_meli.py <pasta de saída>
import random, csv, sys, os
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(20261026)
OUT = sys.argv[1]; os.makedirs(OUT, exist_ok=True)
MESES = {7: 31, 8: 31, 9: 30, 10: 24}
FASE = {7: "Antes do serviço", 8: "Antes do serviço", 9: "Implantação do controle", 10: "Com controle"}
NO_PRAZO = {7: .81, 8: .83, 9: .89, 10: .95}      # pacotes entregues no prazo / entregues
INSUCESSO = {7: .085, 8: .078, 9: .058, 10: .038}  # pacotes não entregues / carregados
SAIDA_ATRASO = {7: .30, 8: .27, 9: .17, 10: .08}   # rotas que saíram depois do horário
BASE = "Base Fictícia · Duque de Caxias/RJ"
ROTAS = {"Duque de Caxias Centro": 18, "São João de Meriti": 22, "Nova Iguaçu": 30, "Belford Roxo": 26, "Nilópolis": 20,
         "Méier / Engenho Novo": 35, "Tijuca": 40, "Barra da Tijuca": 55, "Campo Grande": 60, "Niterói": 45, "São Gonçalo": 50}
VEIC = ["Fiorino", "Van", "VUC"]
NOMES = ["André Silva","Caio Ferreira","Douglas Santos","Elias Costa","Felipe Araújo","Gabriel Melo","Heitor Ribeiro","Ivan Martins",
         "João Pedro Alves","Kleber Souza","Luan Oliveira","Mateus Correia","Natan Barros","Pedro Henrique","Rodrigo Pinto","Samuel Torres"]
MOT = [(n, f"{random.choice(['RIO','KZT','LMB','QRS'])}{random.randint(0,9)}{random.choice('ABCDEFGHJ')}{random.randint(10,99)}") for n in NOMES]
MOT_INS = ["Cliente ausente", "Endereço não localizado", "Área de risco", "Recusa do cliente", "Avaria", "Fora do horário de entrega"]
MOT_ATR = {"Saída": ["Carregamento atrasado na base", "Motorista chegou atrasado", "Roteirização liberada tarde"],
           "Rota": ["Trânsito intenso", "Rota mal sequenciada", "Muitas paradas fora da área"],
           "Cliente": ["Espera no endereço", "Condomínio sem portaria"]}
fmt = lambda d: d.strftime("%d/%m/%Y %H:%M")

cab = ["ID Rota","Mês","Data","Fase","Base","Rota (região)","Distância (km)","Motorista","Placa","Veículo",
       "Saída programada","Saída real","Atraso saída (min)","Retorno programado","Retorno real",
       "Pacotes carregados","Entregues","Entregues no prazo","Insucessos","% no prazo","% insucesso",
       "Principal motivo de insucesso","Motivo do atraso","Detalhe do atraso"]
L = []; n = 90000
for mes, dias in MESES.items():
    for d in range(1, dias + 1):
        for rota, km in ROTAS.items():
            m = random.choice(MOT); carregados = random.randint(55, 120)
            saida_p = datetime(2026, mes, d, 8, random.choice([0, 15, 30]))
            atrasou = random.random() < SAIDA_ATRASO[mes]
            atraso = random.randint(10, 75) if atrasou else 0
            saida_r = saida_p + timedelta(minutes=atraso if atrasou else -random.randint(0, 10))
            ins = max(0, round(carregados * random.gauss(INSUCESSO[mes], 0.02)))
            entregues = carregados - ins
            taxa = min(1, max(0, random.gauss(NO_PRAZO[mes] - (0.06 if atrasou else 0), 0.05)))
            no_prazo = round(entregues * taxa)
            ret_p = saida_p + timedelta(hours=9)
            ret_r = ret_p + timedelta(minutes=random.randint(-40, 20) + (atraso if atrasou else 0) + (0 if taxa > .9 else random.randint(10, 60)))
            if atrasou: grupo = "Saída"
            elif taxa < .85: grupo = random.choice(["Rota", "Rota", "Cliente"])
            else: grupo = ""
            L.append([f"MLR-{n}", mes, saida_p.strftime("%d/%m/%Y"), FASE[mes], BASE, rota, km + random.randint(-3, 8), m[0], m[1], random.choice(VEIC),
                      fmt(saida_p), fmt(saida_r), max(0, atraso), fmt(ret_p), fmt(ret_r), carregados, entregues, no_prazo, ins,
                      round(100 * no_prazo / entregues, 1) if entregues else 0, round(100 * ins / carregados, 1),
                      random.choice(MOT_INS) if ins else "", grupo, random.choice(MOT_ATR[grupo]) if grupo else ""]); n += 1

PLANOS = [
 ("Saída", "Carregamento atrasado na base", "Janela de carregamento por onda, com rota pronta 30 min antes", "Supervisor da base", "set/2026", "Concluído"),
 ("Saída", "Motorista chegou atrasado", "Check-in do motorista às 7h30 com alerta automático aos 10 min de atraso", "Torre de controle", "set/2026", "Concluído"),
 ("Saída", "Roteirização liberada tarde", "Roteirização fechada na noite anterior", "Planejamento", "out/2026", "Em andamento"),
 ("Rota", "Rota mal sequenciada", "Revisar sequência das 3 rotas com pior % no prazo toda segunda", "Analista de performance", "out/2026", "Em andamento"),
 ("Rota", "Trânsito intenso", "Antecipar saída das rotas Barra e Campo Grande em 20 min", "Planejamento", "out/2026", "Concluído"),
 ("Insucesso", "Cliente ausente", "Mensagem ao cliente 1 h antes da entrega", "Atendimento", "out/2026", "Em andamento"),
 ("Insucesso", "Endereço não localizado", "Motorista registra ponto de referência no app; lista de endereços problemáticos", "Torre de controle", "set/2026", "Concluído"),
 ("Insucesso", "Área de risco", "Mapa de áreas de risco com janela de entrega diurna", "Supervisor da base", "nov/2026", "Planejado"),
 ("Insucesso", "Avaria", "Conferência de pacotes frágeis no carregamento", "Supervisor da base", "nov/2026", "Planejado"),
 ("Gestão", "Falta de visibilidade dos números", "Painel diário com % no prazo, insucesso e atraso de saída por rota", "Analista de performance", "set/2026", "Concluído"),
 ("Gestão", "Motoristas sem retorno do desempenho", "Ranking semanal de motoristas enviado no grupo", "Torre de controle", "out/2026", "Em andamento"),
 ("Gestão", "Reincidência das mesmas falhas", "Reunião de 15 min toda sexta com os 3 piores pontos da semana", "Gerente de operações", "out/2026", "Em andamento"),
]
cab_p = ["#", "Tema", "Problema / causa", "Ação", "Responsável", "Prazo", "Status"]

with open(f"{OUT}/meli_rotas_ficticio.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";"); w.writerow(cab); w.writerows(L)
with open(f"{OUT}/meli_planos_de_acao_ficticio.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";"); w.writerow(cab_p); w.writerows([[i + 1, *p] for i, p in enumerate(PLANOS)])

wb = Workbook(); ws = wb.active; ws.title = "LEIA-ME"
for l in ["BASE FICTÍCIA · MODELO MERCADO LIVRE (LAST MILE)", "Dados gerados por código. Não há dado real de nenhuma empresa.",
          "Indicadores genéricos: % no prazo (entregues no prazo ÷ entregues), % insucesso (não entregues ÷ carregados), atraso de saída da base.",
          "Os indicadores oficiais do Mercado Livre devem ser confirmados com o cliente antes de usar em proposta.",
          "Período 01/07 a 24/10/2026 · Fase: jul-ago Antes do serviço · set Implantação do controle · out Com controle.", f"{len(L)} rotas · {len(PLANOS)} planos de ação"]:
    ws.append([l])
ws["A1"].font = Font(bold=True, size=14); ws.column_dimensions["A"].width = 130
head = PatternFill("solid", fgColor="2D3277")
for aba, c, linhas in (("Rotas", cab, L), ("Planos de ação", cab_p, [[i + 1, *p] for i, p in enumerate(PLANOS)])):
    s = wb.create_sheet(aba); s.append(c)
    for x in s[1]: x.font = Font(bold=True, color="FFFFFF"); x.fill = head
    for l in linhas: s.append(l)
    s.freeze_panes = "A2"; s.auto_filter.ref = s.dimensions
    for i, h in enumerate(c, 1): s.column_dimensions[get_column_letter(i)].width = max(10, min(60 if aba != "Rotas" else 30, len(h) + 6))
wb.save(f"{OUT}/base_ficticia_mercado_livre.xlsx")
print("Rotas", len(L), "· Planos", len(PLANOS))
