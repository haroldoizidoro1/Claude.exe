# Gera bases FICTÍCIAS de viagens (Amazon, Shopee, J&T) para painéis de demonstração.
# História: jul-ago sem o serviço (ruim), set = mês da contratação (implantação), out = resultados com controle.
# Uso: python3 gerar_bases.py <pasta de saída>
import random, csv, sys, os
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(20261025)
OUT = sys.argv[1]; os.makedirs(OUT, exist_ok=True)
MESES = {7: 31, 8: 31, 9: 30, 10: 24}  # outubro vai até o dia 24 (mês atual)
FASE = {7: "Antes do serviço", 8: "Antes do serviço", 9: "Implantação do controle", 10: "Com controle"}
NOMES = ["Carlos Menezes","Rafael Souza","Anderson Lima","Bruno Teixeira","Diego Carvalho","Eduardo Pires","Fábio Nunes","Gustavo Rocha",
 "Hugo Barbosa","Igor Cardoso","Jorge Almeida","Leandro Gomes","Marcelo Dias","Nilton Faria","Otávio Ramos","Paulo Moreira",
 "Renato Castro","Sérgio Lopes","Thiago Moura","Vinícius Prado","Wagner Freitas","Alex Monteiro","Cláudio Reis","Danilo Vieira"]
MOT = [(n, f"{random.choice(['RJX','KXA','LTB','QPC','SRD'])}{random.randint(0,9)}{random.choice('ABCDEFGHJ')}{random.randint(10,99)}") for n in NOMES]
VEIC = ["Truck","Carreta","Toco","VUC"]
MOTIVOS = {
 "Origem": ["Carregamento demorado no CD","Motorista apresentou atrasado","Documentação/liberação atrasada","Veículo chegou com pendência"],
 "Trânsito": ["Trânsito intenso","Acidente na rodovia","Parada não programada","Obra na pista","Chuva forte"],
 "Destino": ["Fila de descarga","Base sem doca disponível","Redirecionamento de base","Divergência de lacre na chegada"],
}
fmt = lambda d: d.strftime("%d/%m/%Y %H:%M")
mins = lambda a, b: int(round((a - b).total_seconds() / 60))
M = lambda m: timedelta(minutes=m)

def marca(n, taxa):
    """Lista com exatamente round(n*taxa) True, embaralhada: garante o indicador do mês."""
    k = round(n * taxa); v = [True] * k + [False] * (n - k); random.shuffle(v); return v

def base_viagens(lanes, por_dia):
    out = []
    for mes, dias in MESES.items():
        lst = []
        for d in range(1, dias + 1):
            for _ in range(por_dia + random.randint(-1, 1)):
                lane = random.choice(list(lanes)); tt = lanes[lane]
                std = datetime(2026, mes, d, random.choice([0,2,4,5,6,8,10,13,15,18,20,22]), random.choice([0,15,30,45]))
                lst.append(dict(mes=mes, lane=lane, tt=tt, std=std, sta=std + M(tt), apres_prog=std - M(60), mot=random.choice(MOT)))
        out.append((mes, lst))
    return out

def cli_status(nosso, opcoes, p=0.05):
    return nosso if random.random() > p else random.choice([o for o in opcoes if o != nosso])

def salva(nome, cab, linhas):
    linhas.sort(key=lambda x: datetime.strptime(x[2], "%d/%m/%Y"))
    with open(f"{OUT}/{nome}.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";"); w.writerow(cab); w.writerows(linhas)

# ---------------- AMAZON: atrasos apontados por mês caindo (jul 56 · ago 53 · set 45 · out 26), outubro com 23 de 26 indevidos (88%) ----------------
A_LANES = {"Duque de Caxias/RJ → Cajamar/SP":420,"Duque de Caxias/RJ → Betim/MG":390,"Duque de Caxias/RJ → Serra/ES":480,
 "Nova Iguaçu/RJ → São Gonçalo/RJ":75,"Nova Iguaçu/RJ → Campos dos Goytacazes/RJ":240,"Cajamar/SP → Duque de Caxias/RJ":420,"Betim/MG → Duque de Caxias/RJ":390}
A_ATRASOS = {7: 56, 8: 53, 9: 45, 10: 26}; A_REAL = {7: .88, 8: .87, 9: .65, 10: 3 / 26}
cab_a = ["ID Viagem","Mês","Data","Fase","Lane","Tipo de rota","Motorista","Placa","Veículo","Apresentação programada","Chegada na origem","CPT (STD)","Saída real (ATD)","STA destino","Chegada real (ATA)",
 "Transit time (min)","Atraso chegada origem (min)","Atraso saída (min)","Atraso destino (min)","OTP","OTA","Status nosso","Status Amazon","Divergência","Motivo do atraso","Detalhe do motivo",
 "VRID","Vínculo","ADHOC/Retorno","Operador","Evidência","Amazon OTP","Amazon OTA","Tipo OTP","Tipo OTA","Min. Amazon OTP","Min. Amazon OTA","Resumo OTP","Resumo OTA","Preço BID (R$)"]
# Campos extras com gerador próprio (não mexem nos números já calibrados)
R2 = random.Random(2610)
VR_USADOS = set()
def vrid():
    while True:
        v = "11" + "".join(R2.choice("0123456789ABCDEFGHJKLMNPQRSTVWXYZ") for _ in range(4))
        if v not in VR_USADOS: VR_USADOS.add(v); return v + "***"   # 3 últimos mascarados
OPERADORES = ["Lucas", "Bianca", "Rafael", "Camila"]
RESUMO_REAL = {"Origem": ["Documentação", "Condutor sem evidência", "Quebra mecânica"], "Trânsito": ["Trânsito", "Trânsito", "Erro do condutor em trajeto", "Quebra mecânica", "Saúde do condutor"],
               "Destino": ["Trânsito", "Erro do condutor em trajeto"]}
RESUMO_INDEV = {"OTP": ["Erro da base OTP", "Liberado CPT"], "OTA": ["Erro da base OTA", "Liberado CPT"]}
BID = {"Duque de Caxias/RJ → Cajamar/SP": 4200, "Cajamar/SP → Duque de Caxias/RJ": 4200, "Duque de Caxias/RJ → Betim/MG": 3900, "Betim/MG → Duque de Caxias/RJ": 3900,
       "Duque de Caxias/RJ → Serra/ES": 4600, "Nova Iguaçu/RJ → São Gonçalo/RJ": 950, "Nova Iguaçu/RJ → Campos dos Goytacazes/RJ": 2500}  # preço por carga de CARRETA, fictício
def resumo_ind(t, ind, causa):
    if t == "Indevido": return R2.choice(RESUMO_INDEV[ind])
    if t == "Real": return R2.choice(RESUMO_REAL[causa or "Trânsito"])
    return ""
def extras(lane, veic, mes, a1, a3, amz_o, amz_a):
    min_o = a1 if a1 > 0 else (R2.randint(3, 45) if amz_o else 0)
    min_a = a3 if a3 > 0 else (R2.randint(3, 45) if amz_a else 0)
    tipo = lambda amz, nosso_min: ("Real" if nosso_min > 0 else "Indevido") if amz else ""
    t_o, t_a = tipo(amz_o, a1), tipo(amz_a, a3)
    adhoc = "ADHOC" if ("MG" in lane.split("→")[1] or "ES" in lane.split("→")[1]) else ("Retorno" if lane.endswith("Duque de Caxias/RJ") else "")
    evid = "" if not (amz_o or amz_a) else ("OK" if R2.random() < (.97 if mes == 10 else .85) else "Pendente")
    preco = round(BID[lane] * (1 if veic == "Carreta" else .7))
    return [vrid(), "Agregado" if R2.random() < .7 else "Frota", adhoc, R2.choice(OPERADORES), evid,
            "Atrasado" if amz_o else "No prazo", "Atrasado" if amz_a else "No prazo", t_o, t_a, min_o, min_a, preco], (t_o, t_a)
LA = []; n = 26000
for mes, lst in base_viagens(A_LANES, 11):
    # atrasos apontados pela Amazon no mês (linhas OTP + OTA) e parte que é nossa (Real)
    tot = A_ATRASOS[mes]; reais = round(tot * A_REAL[mes]); indev = tot - reais
    r_otp = round(reais * .6); r_ota = reais - r_otp; i_otp = (indev + 1) // 2; i_ota = indev // 2
    nv = len(lst); ix1 = list(range(nv)); ix2 = list(range(nv)); random.shuffle(ix1); random.shuffle(ix2)
    otpR, otaR = set(ix1[:r_otp]), set(ix2[:r_ota])
    indO = set([i for i in ix2 if i not in otpR][:i_otp]); indA = set([i for i in ix1 if i not in otaR and i not in indO][:i_ota])
    for i, v in enumerate(lst):
        ol, dl = i in otpR, i in otaR
        apres = v["apres_prog"] + (M(random.randint(3, 75)) if ol else -M(random.randint(0, 50)))
        sai_late = (ol and random.random() < .55) or (not ol and random.random() < .01)
        atd = v["std"] + (M(random.randint(3, 60)) if sai_late else -M(random.randint(0, 8)))
        a2 = max(0, mins(atd, v["std"]))
        if dl:
            causa = "Origem" if a2 > 0 and random.random() < .7 else random.choice(["Trânsito", "Trânsito", "Destino"])
            ata = v["sta"] + M(random.randint(4, 120)) if causa != "Origem" else v["sta"] + M(max(1, a2 + random.randint(-5, 30)))
        else:
            ata = v["sta"] - M(random.randint(0, 25)); causa = ""
        a1 = max(0, mins(apres, v["apres_prog"])); a3 = max(0, mins(ata, v["sta"]))
        if ol and not causa: causa = "Origem"
        amz_o, amz_a = ol or i in indO, dl or i in indA
        nosso = "Atrasado" if (a1 > 0 or a3 > 0) else "No prazo"; cli = "Atrasado" if (amz_o or amz_a) else "No prazo"
        LA.append([f"AMZ-{n}", mes, v["std"].strftime("%d/%m/%Y"), FASE[mes], v["lane"], "Ad hoc" if ("MG" in v["lane"] or "ES" in v["lane"]) else "Fixa",
                   v["mot"][0], v["mot"][1], random.choice(VEIC[:2]), fmt(v["apres_prog"]), fmt(apres), fmt(v["std"]), fmt(atd), fmt(v["sta"]), fmt(ata), v["tt"],
                   a1, a2, a3, "Atrasado" if a1 > 0 else "No prazo", "Atrasado" if a3 > 0 else "No prazo", nosso, cli, "Sim" if cli != nosso else "Não",
                   causa, random.choice(MOTIVOS[causa]) if causa else ""]); n += 1
        ex, (t_o, t_a) = extras(v["lane"], LA[-1][8], mes, a1, a3, amz_o, amz_a)
        LA[-1] += ex[:11] + [resumo_ind(t_o, "OTP", causa), resumo_ind(t_a, "OTA", causa), ex[11]]
salva("amazon_viagens_ficticio", cab_a, LA)

# ---------------- SHOPEE: começa entre 60% e 68% na origem e no destino, sobe a partir de setembro ----------------
S_LANES = {"Duque de Caxias/RJ → São Gonçalo/RJ":70,"Duque de Caxias/RJ → Niterói/RJ":55,"São João de Meriti/RJ → Campo Grande/RJ":95,
 "Duque de Caxias/RJ → Petrópolis/RJ":110,"Duque de Caxias/RJ → Volta Redonda/RJ":150,"Duque de Caxias/RJ → Macaé/RJ":230,"São João de Meriti/RJ → Cabo Frio/RJ":190}
S_ORIG = {7: .63, 8: .66, 9: .76, 10: .88}; S_DEST = {7: .61, 8: .64, 9: .74, 10: .86}
def tol_shopee(tt):
    for lim, t in [(60,5),(180,10),(360,15),(720,20),(1440,25)]:
        if tt < lim: return t
    return 30
cab_s = ["ID Viagem","Mês","Data","Fase","Lane","Motorista","Placa","Veículo","Pacotes","Chegada origem programada","Chegada na origem","STD (CPT)","Saída real (ATD)","STA destino","Chegada real (ATA)",
 "Transit time (min)","Tolerância destino (min)","Ajuste de saída (min)","STA ajustado","Janela início","Janela fim","Atraso origem (min)","Resultado origem","Resultado destino",
 "Status Shopee","Divergência","OCC","Motivo do atraso","Detalhe do motivo"]
LS = []; n = 48000
for mes, lst in base_viagens(S_LANES, 10):
    o_bad = marca(len(lst), 1 - S_ORIG[mes]); d_bad = marca(len(lst), 1 - S_DEST[mes])
    for v, ob, db in zip(lst, o_bad, d_bad):
        tol = tol_shopee(v["tt"])
        apres = v["apres_prog"] + (M(random.randint(3, 70)) if ob else -M(random.randint(0, 45)))  # adiantado na origem não pontua
        r = random.random()
        ajuste = random.randint(5, 50) if (ob and r < .5) else (-random.randint(10, 40) if r > .88 else random.randint(-3, 3))
        atd = v["std"] + M(ajuste); sta_aj = v["sta"] + M(ajuste); ini, fim = sta_aj - M(tol), sta_aj + M(tol)
        if db:
            early = random.random() < .35
            ata = ini - M(random.randint(1, 40)) if early else fim + M(random.randint(1, 110))
        else:
            ata = sta_aj + M(random.randint(-tol, tol))
        atr_o = max(0, mins(apres, v["apres_prog"]))
        r_orig = "Delay" if atr_o > 0 else "No prazo"
        r_dest = "Early" if ata < ini else ("Delay" if ata > fim else "No prazo")
        occ = "Sim" if r_dest != "No prazo" and random.random() < .06 else "Não"
        cli = cli_status(r_dest, ["Early", "Delay", "No prazo"], .06)
        if r_dest == "Delay": causa = "Origem" if ajuste > 0 and random.random() < .6 else random.choice(["Trânsito", "Trânsito", "Destino"])
        elif r_dest == "Early": causa = "Destino"
        elif r_orig == "Delay": causa = "Origem"
        else: causa = ""
        det = "Chegou antes da janela (Early)" if r_dest == "Early" else (random.choice(MOTIVOS[causa]) if causa else "")
        LS.append([f"SPX-{n}", mes, v["std"].strftime("%d/%m/%Y"), FASE[mes], v["lane"], v["mot"][0], v["mot"][1], random.choice(VEIC), random.randint(380, 2600),
                   fmt(v["apres_prog"]), fmt(apres), fmt(v["std"]), fmt(atd), fmt(v["sta"]), fmt(ata), v["tt"], tol, ajuste, fmt(sta_aj), fmt(ini), fmt(fim),
                   atr_o, r_orig, r_dest, cli, "Sim" if cli != r_dest else "Não", occ, causa, det]); n += 1
salva("shopee_viagens_ficticio", cab_s, LS)

# ---------------- J&T: dívida de ~R$ 15 mil por falta de app em julho, caindo até outubro ----------------
J_LANES = {"Duque de Caxias/RJ → Nova Iguaçu/RJ":50,"Duque de Caxias/RJ → Niterói/RJ":60,"Duque de Caxias/RJ → Itaboraí/RJ":85,"São João de Meriti/RJ → Resende/RJ":170,
 "Duque de Caxias/RJ → Teresópolis/RJ":120,"Duque de Caxias/RJ → Angra dos Reis/RJ":200}
J_SEM_APP = {7: (.38, .32), 8: (.26, .28), 9: (.12, .18), 10: (.03, .07)}  # (não usou nos dois pontos, não usou em um ponto)
J_ATRASO = {7: .08, 8: .07, 9: .04, 10: .02}
cab_j = ["ID Viagem","Mês","Data","Fase","Lane","Motorista","Placa","Veículo","CPT (STD)","Saída real (ATD)","STA destino","Chegada real (ATA)","Transit time (min)","Atraso destino (min)",
 "Usou app na saída","Usou app na chegada","Uso do app (%)","Desconto atraso (R$)","Desconto app (R$)","Desconto total nosso (R$)","Desconto lançado pela J&T (R$)","Divergência","Motivo do atraso","Detalhe do motivo"]
LJ = []; n = 71000
for mes, lst in base_viagens(J_LANES, 9):
    p2, p1 = J_SEM_APP[mes]; k2, k1 = round(len(lst) * p2), round(len(lst) * p1)
    app = ["dois"] * k2 + ["um"] * k1 + ["ok"] * (len(lst) - k2 - k1); random.shuffle(app)
    atr = marca(len(lst), J_ATRASO[mes])
    for v, ap, at in zip(lst, app, atr):
        atd = v["std"] + (M(random.randint(2, 30)) if at and random.random() < .4 else -M(random.randint(0, 6)))
        ata = v["sta"] + M(random.randint(2, 45)) if at else v["sta"] - M(random.randint(0, 20))
        a = max(0, mins(ata, v["sta"]))  # tolerância ZERO
        if ap == "dois": s, c_ = "Não", "Não"
        elif ap == "um": s, c_ = random.choice([("Não", "Sim"), ("Sim", "Não")])
        else: s, c_ = "Sim", "Sim"
        uso = 100 * ((s == "Sim") + (c_ == "Sim")) // 2
        d_atr = a * 10; d_app = {"dois": 100, "um": 50, "ok": 0}[ap]; tot = d_atr + d_app
        jt = max(0, tot if random.random() > .05 else tot + random.choice([10, 20, 50, -10, 30]))
        causa = ("Origem" if mins(atd, v["std"]) > 0 else random.choice(["Trânsito", "Trânsito", "Destino"])) if a > 0 else ""
        LJ.append([f"JT-{n}", mes, v["std"].strftime("%d/%m/%Y"), FASE[mes], v["lane"], v["mot"][0], v["mot"][1], random.choice(VEIC[1:]), fmt(v["std"]), fmt(atd), fmt(v["sta"]), fmt(ata),
                   v["tt"], a, s, c_, uso, d_atr, d_app, tot, jt, "Sim" if jt != tot else "Não", causa, random.choice(MOTIVOS[causa]) if causa else ""]); n += 1
salva("jt_viagens_ficticio", cab_j, LJ)

# ---------------- Excel único ----------------
REGRAS = {
 "Amazon": ["Tolerância ZERO nos três pontos: chegada na origem (OTP), saída (CPT) e chegada no destino (OTA).","VRID fictício com os 3 últimos caracteres mascarados (***). Tipo Real = atraso nosso; Indevido = a Amazon apontou, mas nosso registro está no horário.",
            "Preço BID fictício por carga (Carreta; Truck = 70%). As regras de faturamento ficam no painel e são editáveis.",
            "Status Amazon = número do cliente; Status nosso = calculado por nós. Divergência quando diferem."],
 "Shopee": ["Origem: SEM tolerância. Chegou atrasado = Delay. Adiantado na origem não pontua.",
            "Destino: janela STA ajustado ± tolerância pelo transit time (até 1h: 5 · 1h–2h59: 10 · 3h–5h59: 15 · 6h–11h59: 20 · 12h–23h59: 25 · 24h+: 30 min).",
            "Ajuste de saída: liberação antecipada (negativo) puxa o STA para antes; morosidade no carregamento (positivo) empurra para depois.",
            "Early e Delay só pontuam no destino. OCC = ocorrência anulada pela Shopee. Contrato/SLA completo: PENDENTE."],
 "J&T": ["Tolerância ZERO. R$ 10 por minuto de atraso na chegada.","Uso obrigatório do app: R$ 50 se faltar em um ponto (saída ou chegada), R$ 100 se faltar nos dois. Acumula por viagem.",
         "Uso do app (%) = pontos registrados no app ÷ 2. Desconto lançado pela J&T = número do cliente."],
}
wb = Workbook(); ws = wb.active; ws.title = "LEIA-ME"
for l in ["BASE FICTÍCIA PARA DEMONSTRAÇÃO",
          "Todos os dados (motoristas, placas, IDs, horários, descontos) foram gerados por código. Não há dado real de nenhuma empresa.",
          "Período: 01/07/2026 a 24/10/2026. Coluna Fase: jul-ago = Antes do serviço · set = Implantação do controle · out = Com controle.", ""]:
    ws.append([l])
ws["A1"].font = Font(bold=True, size=14)
for (aba, linhas) in (("Amazon", LA), ("Shopee", LS), ("J&T", LJ)):
    ws.append([f"{aba}: {len(linhas)} viagens"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    for r in REGRAS[aba]: ws.append(["  • " + r])
    ws.append([])
ws.column_dimensions["A"].width = 150
head = PatternFill("solid", fgColor="2A322E")
for aba, cab, linhas in (("Amazon", cab_a, LA), ("Shopee", cab_s, LS), ("J&T", cab_j, LJ)):
    s = wb.create_sheet(aba); s.append(cab)
    for c in s[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = head
    for l in linhas: s.append(l)
    s.freeze_panes = "A2"; s.auto_filter.ref = s.dimensions
    for i, h in enumerate(cab, 1): s.column_dimensions[get_column_letter(i)].width = max(10, min(34, len(h) + 4))
wb.save(f"{OUT}/base_ficticia_3_clientes.xlsx")
print("Amazon", len(LA), "· Shopee", len(LS), "· J&T", len(LJ))
